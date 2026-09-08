#!/usr/bin/env python3
"""Rebuild the Canonical Dashboard sheet from the Canonical Backlog sheet.

WHY THIS EXISTS. The dashboard is a summary of the backlog and had no way to
stay one. It was recomputed by hand, twice, and drifted both times. On 6 Sept
2026 it read 124 issues against a sheet holding 156. It was corrected on 7 Sept
- but only the headline tiles and the Status table were corrected, so the Issue
Type, Risk Area, Canonical Theme and Target Release blocks were left at their
124-row state and were still wrong the next morning, when the Status table had
drifted again as well. A number in a spreadsheet is a fact with no owner and
nothing fails when it goes stale, which is the same reason CLAUDE.md forbids
writing counts into it. So the dashboard is generated now, and this is the
generator. Run it after any change to the backlog.

WHAT IT COUNTS, exactly. Every number names the column it comes from, on the
sheet itself, because until now nobody could tell which column drove a tile -
"P1 / High Priority" read 62 and no column in the workbook produced 62.

  Total Issues          rows in Canonical Backlog with a non-empty Issue ID
  P1 / High Priority    Priority == "P1 - High"          (all rows)
  Decision Needed       Decision Needed? == "Yes"        (all rows)
  Needs Clarification   Status == "Needs Clarification"
  Blocked               Status == "Blocked"
  Ready for Build/QA    Status == "Ready for QA"
  each summary block    one row per distinct value of that column, blanks
                        counted separately and shown as "(not set)"

The three tiles that are not status-based count ALL rows, the same population
as Total Issues, because their labels carry no qualifier. The open-only figures
are printed beside them in the definitions block, since "72 P1 issues, 7 of
them still open" is the pair a reader actually wants.

Usage:
    python3 scripts/regenerate_dashboard.py [workbook.xlsx] [--dry-run]

With no path it takes the newest CAO_OMS_Canonical_Backlog_CURRENT_*.xlsx in
the Desktop OMS folder. --dry-run prints what would change and writes nothing.
Back the workbook up before a real run; the convention in that folder is
CAO_OMS_Canonical_Backlog_BACKUP_before_<reason>.xlsx.
"""
import collections
import datetime
import glob
import os
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.stderr.write("openpyxl is required: pip install openpyxl\n")
    raise SystemExit(2)

DEFAULT_GLOB = os.path.join(
    os.path.expanduser("~"), "Desktop", "Work", "OMS",
    "CAO_OMS_Canonical_Backlog_CURRENT_*.xlsx")

BACKLOG = "Canonical Backlog"
DASHBOARD = "Canonical Dashboard"

NAVY = "FF1F4E78"
PALE = "FFD9EAF7"
WHITE = "FFFFFFFF"

# Blocks are (heading, backlog column). Order is the order they appear.
BLOCKS = [
    ("Status Summary", "Status"),
    ("Issue Type Summary", "Issue Type"),
    ("Risk Area Summary", "Risk Area"),
    ("Canonical Theme Summary", "Canonical Theme"),
    ("Target Release Summary", "Target Release"),
]

# A row counts as still open unless its Status says otherwise. Used only for
# the open-only figures in the definitions block, never for a tile.
CLOSED_STATUSES = {"closed", "closed - documented limitation", "won't do",
                   "wont do", "deferred"}


def norm(v):
    return str(v).strip() if v is not None else ""


def read_backlog(ws):
    hdr = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}
    if "Issue ID" not in hdr:
        raise SystemExit("'%s' has no Issue ID column - wrong sheet?" % ws.title)
    rows = []
    for r in range(2, ws.max_row + 1):
        if not norm(ws.cell(r, hdr["Issue ID"]).value):
            continue
        rows.append({h: ws.cell(r, i).value for h, i in hdr.items()})
    return hdr, rows


def tally(rows, column):
    """Counts by value, commonest first, blanks last under '(not set)'."""
    counts = collections.Counter(norm(x.get(column)) for x in rows)
    blank = counts.pop("", 0)
    out = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0].lower()))
    if blank:
        out.append(("(not set)", blank))
    return out


def read_existing(ws, tile_labels):
    """Whatever numbers are on the sheet now, so the run can show what moved.

    Keyed by (block heading, label) rather than by label alone. A flat dict
    collides across blocks - "Data Integrity" is a Risk Area with dozens of
    rows AND a Canonical Theme with a couple - and reports both against
    whichever it saw first, which is a diff that invents changes.
    """
    seen = {}
    # Tiles: label on one row, value directly beneath it in the same column.
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip() in tile_labels:
                below = ws.cell(c.row + 1, c.column).value
                if isinstance(below, int):
                    seen[("tile", c.value.strip())] = below
    # Summary blocks: a heading in column A with "Count" beside it, then rows
    # of label and count until the first gap.
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not (isinstance(a, str) and isinstance(b, str) and b.strip().lower() == "count"):
            continue
        heading = a.strip()
        rr = r + 1
        while rr <= ws.max_row:
            label, count = ws.cell(rr, 1).value, ws.cell(rr, 2).value
            if not isinstance(label, str) or not isinstance(count, int):
                break
            seen[(heading, label.strip())] = count
            rr += 1
    return seen


def build(rows, hdr):
    n = len(rows)
    is_open = [x for x in rows
               if norm(x.get("Status")).lower() not in CLOSED_STATUSES]

    def count(col, value):
        return sum(1 for x in rows if norm(x.get(col)) == value)

    tiles = [
        ("Total Issues", n, "rows with an Issue ID"),
        ("P1 / High Priority", count("Priority", "P1 - High"), 'Priority == "P1 - High"'),
        ("Decision Needed", count("Decision Needed?", "Yes"), 'Decision Needed? == "Yes"'),
        ("Needs Clarification", count("Status", "Needs Clarification"), 'Status == "Needs Clarification"'),
        ("Blocked", count("Status", "Blocked"), 'Status == "Blocked"'),
        ("Ready for Build/QA", count("Status", "Ready for QA"), 'Status == "Ready for QA"'),
    ]
    open_p1 = sum(1 for x in is_open if norm(x.get("Priority")) == "P1 - High")
    open_dec = sum(1 for x in is_open if norm(x.get("Decision Needed?")) == "Yes")
    blocks = [(title, col, tally(rows, col)) for title, col in BLOCKS if col in hdr]
    return tiles, blocks, len(is_open), open_p1, open_dec


def write(wb, tiles, blocks, n_open, open_p1, open_dec, stamp):
    idx = wb.sheetnames.index(DASHBOARD) if DASHBOARD in wb.sheetnames else 0
    if DASHBOARD in wb.sheetnames:
        del wb[DASHBOARD]
    ws = wb.create_sheet(DASHBOARD, idx)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26.0
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 13.0

    head_font = Font(size=11, bold=True, color=WHITE)
    navy = PatternFill("solid", fgColor=NAVY)
    pale = PatternFill("solid", fgColor=PALE)
    centre = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")

    def merge(row, first, last):
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)

    ws["A1"] = "CAO OMS Canonical Feature Request Backlog Dashboard"
    ws["A1"].font = Font(size=16, bold=True, color=WHITE)
    ws["A1"].fill = navy
    merge(1, 1, 8)

    ws["A2"] = ("Generated %s by scripts/regenerate_dashboard.py from the Canonical "
                "Backlog sheet. Do not type a number into this sheet: run the script "
                "again instead. Every count names the column it comes from in the "
                "definitions below." % stamp)
    ws["A2"].alignment = wrap
    merge(2, 1, 8)
    ws.row_dimensions[2].height = 30

    # Two bands of three tiles, at columns A, D and G.
    for band in (0, 1):
        lrow, vrow = 4 + band * 3, 5 + band * 3
        for slot in range(3):
            label, value, _ = tiles[band * 3 + slot]
            col = 1 + slot * 3
            lc, vc = ws.cell(lrow, col), ws.cell(vrow, col)
            lc.value, lc.font, lc.fill, lc.alignment = label, head_font, navy, centre
            vc.value, vc.font, vc.fill, vc.alignment = value, Font(size=14, bold=True), pale, centre
            merge(lrow, col, col + 1)
            merge(vrow, col, col + 1)

    row = 11
    ws.cell(row, 1, "How every number on this sheet is counted").font = Font(bold=True)
    merge(row, 1, 8)
    row += 1
    for label, value, rule in tiles:
        ws.cell(row, 1, "%s: %s" % (label, rule)).alignment = wrap
        merge(row, 1, 8)
        row += 1
    ws.cell(row, 1,
            "Tiles count every row, the same population as Total Issues, because their "
            "labels carry no qualifier. Of the %d issues, %d are still open (Status is "
            "not Closed, Closed - documented limitation, Won't Do or Deferred); %d of "
            "the P1 - High issues and %d of the rows marked Decision Needed are among "
            "them." % (tiles[0][1], n_open, open_p1, open_dec)).alignment = wrap
    merge(row, 1, 8)
    ws.row_dimensions[row].height = 30
    row += 1
    ws.cell(row, 1,
            "Each summary below lists one row per distinct value of that column and "
            "sums to the total issue count, with blanks shown as (not set).").alignment = wrap
    merge(row, 1, 8)
    row += 2

    for title, column, pairs in blocks:
        h1, h2 = ws.cell(row, 1, title), ws.cell(row, 2, "Count")
        for c in (h1, h2):
            c.font, c.fill = head_font, navy
        ws.cell(row, 3, "from column: %s" % column).font = Font(italic=True, size=9)
        row += 1
        for label, count in pairs:
            ws.cell(row, 1, label)
            vc = ws.cell(row, 2, count)
            vc.alignment = centre
            row += 1
        total = ws.cell(row, 1, "Total")
        tv = ws.cell(row, 2, sum(c for _, c in pairs))
        total.font = tv.font = Font(bold=True)
        tv.alignment = centre
        row += 2
    return ws


def main(argv):
    args = [a for a in argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in argv
    path = args[0] if args else None
    if not path:
        found = sorted(glob.glob(DEFAULT_GLOB))
        if not found:
            raise SystemExit("no workbook given and none matched %s" % DEFAULT_GLOB)
        path = found[-1]
    print("workbook: %s%s" % (path, "   (DRY RUN)" if dry else ""))

    wb = openpyxl.load_workbook(path)
    if BACKLOG not in wb.sheetnames:
        raise SystemExit("no '%s' sheet in %s" % (BACKLOG, path))
    hdr, rows = read_backlog(wb[BACKLOG])
    tiles, blocks, n_open, open_p1, open_dec = build(rows, hdr)
    before = (read_existing(wb[DASHBOARD], {t[0] for t in tiles})
              if DASHBOARD in wb.sheetnames else {})
    stamp = datetime.date.today().isoformat()

    print("\n%-34s %8s %8s" % ("", "was", "now"))
    changed = 0
    for label, value, _ in tiles:
        was = before.get(("tile", label))
        flag = "" if was == value else "   <-- changed"
        changed += 0 if was == value else 1
        print("%-34s %8s %8s%s" % (label, "-" if was is None else was, value, flag))
    for title, column, pairs in blocks:
        print("\n  %s (from %s), sums to %d" % (title, column, sum(c for _, c in pairs)))
        for label, count in pairs:
            was = before.get((title, label))
            if was != count:
                changed += 1
                print("      %-44s %s -> %s" % (label[:44], "-" if was is None else was, count))
    print("\n%d number(s) differ from what is on the sheet." % changed)

    if dry:
        print("dry run: nothing written.")
        return 0
    write(wb, tiles, blocks, n_open, open_p1, open_dec, stamp)
    wb.save(path)
    print("written, and stamped %s." % stamp)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
